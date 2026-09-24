import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# INPUT FILES
# ============================================================

FILE_S9A = "Supplementary_Table_S9A.csv"
FILE_S9B = "Supplementary_Table_S9B.csv"
FILE_S9C = "Supplementary_Table_S9C.csv"

OUTPUT_FILE = "Supplementary_Table_S9_WORD.docx"


# ============================================================
# READ AUDITED SOURCE FILES
# ============================================================

s9a = pd.read_csv(FILE_S9A)
s9b = pd.read_csv(FILE_S9B)
s9c = pd.read_csv(FILE_S9C)


# ============================================================
# VALIDATION
# ============================================================

assert len(s9a) == 7, f"Unexpected S9A row count: {len(s9a)}"
assert len(s9b) == 7, f"Unexpected S9B row count: {len(s9b)}"
assert len(s9c) == 10, f"Unexpected S9C row count: {len(s9c)}"

assert set(s9a["Cell_line"]) == {"A549", "PC3", "VCAP"}
assert set(s9b["Cell_line"]) == {"A549", "PC3", "VCAP"}
assert set(s9c["Gene_space"]) == {"Full", "Landmark"}

expected_hallmarks = {
    "HALLMARK_ESTROGEN_RESPONSE_EARLY",
    "HALLMARK_ESTROGEN_RESPONSE_LATE",
    "HALLMARK_E2F_TARGETS",
    "HALLMARK_G2M_CHECKPOINT",
    "HALLMARK_MYC_TARGETS_V1",
}

assert set(s9c["Hallmark_programme"]) == expected_hallmarks


# ============================================================
# PUBLICATION COPIES
# Source CSV files remain unchanged.
# ============================================================

a_pub = s9a.copy()
b_pub = s9b.copy()
c_pub = s9c.copy()

a_pub = a_pub.rename(columns={
    "Cell_line": "Cell line",
    "Experimental_block": "Experimental block",
    "N_6h_profiles": "6 h profiles",
    "N_24h_profiles": "24 h profiles",
    "Has_6h": "6 h present",
    "Has_24h": "24 h present",
    "Matched_6h_24h": "Matched 6–24 h",
})

b_pub = b_pub.rename(columns={
    "Cell_line": "Cell line",
    "Experimental_block": "Experimental block",
    "Pearson_6h_vs_24h": "Pearson r",
    "Spearman_6h_vs_24h": "Spearman rho",
    "Cosine_6h_vs_24h": "Cosine similarity",
    "Mean_absolute_difference": "Mean absolute difference",
    "Mean_absolute_MODZ_6h": "Mean |MODZ| 6 h",
    "Mean_absolute_MODZ_24h": "Mean |MODZ| 24 h",
    "Delta_24h_minus_6h": "Delta |MODZ| (24 h − 6 h)",
})

c_pub = c_pub.rename(columns={
    "Gene_space": "Gene space",
    "Hallmark_programme": "Hallmark programme",
    "N_genes": "Genes",
    "N_Landmark_genes": "Landmark genes",
    "Mean_6h": "Mean 6 h",
    "Mean_24h": "Mean 24 h",
    "Median_6h": "Median 6 h",
    "Median_24h": "Median 24 h",
    "Mean_delta_24h_minus_6h": "Mean delta (24 h − 6 h)",
    "Median_delta_24h_minus_6h": "Median delta (24 h − 6 h)",
    "N_delta_negative": "Negative deltas",
    "N_delta_positive": "Positive deltas",
    "N_negative_6h": "Negative at 6 h",
    "N_positive_6h": "Positive at 6 h",
    "N_negative_24h": "Negative at 24 h",
    "N_positive_24h": "Positive at 24 h",
    "Wilcoxon_statistic": "Wilcoxon statistic",
    "P_two_sided": "Two-sided P",
    "Q_BH": "BH-adjusted P",
})


# ============================================================
# DISPLAY FUNCTIONS
# ============================================================

def display_value(value):
    if pd.isna(value):
        return ""

    if isinstance(value, bool):
        return "Yes" if value else "No"

    if isinstance(value, (int,)):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        if abs(value) < 0.001 and value != 0:
            return f"{value:.3e}"
        return f"{value:.4f}"

    text = str(value)

    if text == "True":
        return "Yes"
    if text == "False":
        return "No"

    return text


# ============================================================
# WORD HELPERS
# ============================================================

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_cell_margins(cell, top=45, start=45, bottom=45, end=45):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)

    for margin, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_cm):
    for row in table.rows:
        for idx, width_cm in enumerate(widths_cm):
            if idx >= len(row.cells):
                continue

            cell = row.cells[idx]
            tcPr = cell._tc.get_or_add_tcPr()

            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)

            tcW.set(qn("w:w"), str(int(width_cm * 567)))
            tcW.set(qn("w:type"), "dxa")


def add_dataframe_table(document, dataframe, widths_cm=None, font_size=7):
    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    header = table.rows[0]
    set_repeat_table_header(header)

    for j, column in enumerate(dataframe.columns):
        cell = header.cells[j]
        cell.text = str(column)
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)

        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(font_size)

    for _, data_row in dataframe.iterrows():
        cells = table.add_row().cells

        for j, value in enumerate(data_row):
            cells[j].text = display_value(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[j])

            for paragraph in cells[j].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)

    if widths_cm is not None:
        set_table_widths(table, widths_cm)

    return table


def add_section_heading(document, text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)

    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)


# ============================================================
# CREATE WORD DOCUMENT
# ============================================================

doc = Document()

section = doc.sections[0]
section.orientation = 1
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(1.3)
section.bottom_margin = Cm(1.3)
section.left_margin = Cm(1.2)
section.right_margin = Cm(1.2)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(8)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(7)

run = title.add_run(
    "Supplementary Table S9. Temporal sensitivity analysis of estradiol "
    "transcriptional responses at 6 and 24 h."
)
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(10)


# ============================================================
# S9A
# ============================================================

add_section_heading(
    doc,
    "Table S9A. Matched experimental blocks included in the temporal analysis."
)

add_dataframe_table(
    doc,
    a_pub,
    widths_cm=[2.2, 3.0, 2.4, 2.4, 2.2, 2.2, 2.8],
    font_size=7.5,
)


# ============================================================
# S9B
# ============================================================

doc.add_page_break()

add_section_heading(
    doc,
    "Table S9B. Global transcriptional similarity and response magnitude "
    "between 6 and 24 h."
)

add_dataframe_table(
    doc,
    b_pub,
    widths_cm=[1.7, 2.3, 2.0, 2.1, 2.2, 3.0, 2.5, 2.5, 3.0],
    font_size=7,
)


# ============================================================
# S9C
# ============================================================

doc.add_page_break()

add_section_heading(
    doc,
    "Table S9C. Temporal changes in selected Hallmark programme scores."
)

add_dataframe_table(
    doc,
    c_pub,
    widths_cm=[
        1.5, 4.0, 1.1, 1.4,
        1.3, 1.3, 1.3, 1.3,
        1.8, 1.8,
        1.3, 1.3,
        1.3, 1.3, 1.3, 1.3,
        1.6, 1.5, 1.7
    ],
    font_size=6.2,
)


# ============================================================
# NOTE
# ============================================================

doc.add_paragraph()

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(3)
note.paragraph_format.space_after = Pt(0)

note_run = note.add_run(
    "Note: Temporal sensitivity of the estradiol transcriptional response was "
    "evaluated using seven matched 6-to-24-h experimental blocks across three "
    "cell lines (A549, PC3 and VCAP). Table S9A reports the matched experimental "
    "blocks included in the analysis. Table S9B summarises within-block global "
    "transcriptional similarity between 6 and 24 h using Pearson correlation, "
    "Spearman rank correlation and cosine similarity, together with measures of "
    "transcriptional-response magnitude. Table S9C reports temporal changes in "
    "mean MODZ scores for HALLMARK_ESTROGEN_RESPONSE_EARLY, "
    "HALLMARK_ESTROGEN_RESPONSE_LATE, HALLMARK_E2F_TARGETS, "
    "HALLMARK_G2M_CHECKPOINT and HALLMARK_MYC_TARGETS_V1 in the Full and "
    "Landmark gene spaces. Paired 6-to-24-h differences were evaluated using "
    "two-sided Wilcoxon signed-rank tests, with Benjamini–Hochberg correction "
    "across the five Hallmark programmes within each gene space. The seven "
    "matched blocks originated from three cell lines and were therefore not "
    "interpreted as seven independent cellular contexts. The Landmark analysis "
    "was considered a robustness analysis rather than an independent validation."
)

note_run.font.name = "Arial"
note_run.font.size = Pt(7.5)


# ============================================================
# SAVE
# ============================================================

doc.save(OUTPUT_FILE)

print("S9 WORD VERSION CREATED CORRECTLY")
print("---------------------------------")
print(f"Output file: {OUTPUT_FILE}")
print(f"S9A rows: {len(s9a)}")
print(f"S9B rows: {len(s9b)}")
print(f"S9C rows: {len(s9c)}")
print()
print("The audited S9A, S9B and S9C CSV files were NOT modified.")