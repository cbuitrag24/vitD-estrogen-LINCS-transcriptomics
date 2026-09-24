import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FILE_A = "Supplementary_Table_S8A.csv"
FILE_B = "Supplementary_Table_S8B.csv"
FILE_C = "Supplementary_Table_S8C.csv"
OUTPUT_FILE = "Supplementary_Table_S8_WORD.docx"

TITLE = (
    "Supplementary Table S8. Perturbation-associated transcriptional responses "
    "of vitamin D- and oestrogen-related receptor-axis genes."
)

NOTE = (
    "Note: Nine genes related to vitamin D and oestrogen signalling were predefined "
    "for analysis: VDR, RXRA, RXRB, RXRG, CYP24A1 and CYP27B1 for the vitamin "
    "D-related axis, and ESR1, ESR2 and GPER1 for the oestrogen-related axis. "
    "Table S8A reports their L1000 measurement status; GPER1 was directly measured "
    "as a Landmark transcript, whereas the remaining eight genes were computationally "
    "inferred. Table S8B summarises treatment-associated MODZ responses and their "
    "directional consistency across the four matched experimental contexts. Table S8C "
    "reports context-specific responses to estradiol, calcitriol and calcipotriol, "
    "together with the mean response to the two vitamin D perturbations and the "
    "difference between the estradiol response and this vitamin D mean. These analyses "
    "were descriptive and no inferential statistical testing was performed. Level 5 "
    "MODZ values represent differential transcriptional responses to perturbation and "
    "should not be interpreted as measures of basal gene expression or direct receptor "
    "functional activity."
)

# Read audited files
a = pd.read_csv(FILE_A)
b = pd.read_csv(FILE_B)
c = pd.read_csv(FILE_C)

# Safety checks
if len(a) != 9:
    raise ValueError(f"Unexpected S8A row count: {len(a)}. Expected 9.")

if len(b) != 27:
    raise ValueError(f"Unexpected S8B row count: {len(b)}. Expected 27.")

if len(c) != 36:
    raise ValueError(f"Unexpected S8C row count: {len(c)}. Expected 36.")

# Publication-friendly headings
a_pub = a.rename(columns={
    "Gene": "Gene",
    "Receptor_axis": "Receptor axis",
    "Present_in_L1000": "Present in L1000",
    "Landmark_status": "Landmark status",
    "L1000_measurement_type": "L1000 measurement type",
}).copy()

b_pub = b.rename(columns={
    "Treatment": "Treatment",
    "Receptor_axis": "Receptor axis",
    "Gene": "Gene",
    "Mean_MODZ": "Mean MODZ",
    "Median_MODZ": "Median MODZ",
    "Min_MODZ": "Minimum MODZ",
    "Max_MODZ": "Maximum MODZ",
    "N_contexts": "Contexts",
    "N_negative": "Negative",
    "N_positive": "Positive",
    "Context_consistency": "Context consistency",
    "L1000_measurement_type": "L1000 measurement type",
}).copy()

c_pub = c.rename(columns={
    "Gene": "Gene",
    "Receptor_axis": "Receptor axis",
    "Cell_context": "Cell context",
    "Estradiol_MODZ": "Estradiol MODZ",
    "Calcitriol_MODZ": "Calcitriol MODZ",
    "Calcipotriol_MODZ": "Calcipotriol MODZ",
    "Mean_vitamin_D_MODZ": "Mean vitamin D MODZ",
    "Estradiol_minus_vitamin_D": "Estradiol − vitamin D",
    "L1000_measurement_type": "L1000 measurement type",
}).copy()

# Reader-friendly Yes/No labels in Word copy only
for column in ["Present in L1000", "Landmark status"]:
    if column in a_pub.columns:
        a_pub[column] = a_pub[column].replace({
            True: "Yes",
            False: "No",
            1: "Yes",
            0: "No",
            "1": "Yes",
            "0": "No",
        })

# Editorial translations in Word copy only
translation_map = {
    "heterogeneo": "Heterogeneous",
    "negativo_4de4": "Negative (4/4)",
    "positivo_4de4": "Positive (4/4)",
    "predominio_negativo_3de4": "Predominantly negative (3/4)",
    "predominio_positivo_3de4": "Predominantly positive (3/4)",
    "Oestrogen": "Estrogen",
    "Vitamin D": "Vitamin D",
    "Inferred": "Inferred",
    "Landmark_direct": "Landmark (direct)",
}

for dataframe in [a_pub, b_pub, c_pub]:
    for column in dataframe.columns:
        if pd.api.types.is_string_dtype(dataframe[column]):
            dataframe[column] = dataframe[column].replace(translation_map)

def display_value(value):
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        return f"{value:.4f}"
    return str(value)


def shade(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9E2F3")
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9)


def add_table(doc, dataframe):
    table = doc.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"
    table.autofit = True

    repeat_header(table.rows[0])

    for j, col in enumerate(dataframe.columns):
        cell = table.rows[0].cells[j]
        cell.text = str(col)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell)

        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(7)

    for _, row in dataframe.iterrows():
        cells = table.add_row().cells

        for j, value in enumerate(row):
            cells[j].text = display_value(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for paragraph in cells[j].paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7)


# Create Word document
doc = Document()

section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(1.3)
section.bottom_margin = Cm(1.3)
section.left_margin = Cm(1.3)
section.right_margin = Cm(1.3)

doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(8)

# Title
p = doc.add_paragraph()
r = p.add_run(TITLE)
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(10)

# S8A
add_heading(
    doc,
    "S8A. L1000 measurement status of receptor-axis genes"
)
add_table(doc, a_pub)

doc.add_page_break()

# S8B
add_heading(
    doc,
    "S8B. Treatment-associated receptor-axis transcriptional responses"
)
add_table(doc, b_pub)

doc.add_page_break()

# S8C
add_heading(
    doc,
    "S8C. Context-specific receptor-axis transcriptional responses"
)
add_table(doc, c_pub)

# Note
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
r = p.add_run(NOTE)
r.font.name = "Arial"
r.font.size = Pt(8)

doc.save(OUTPUT_FILE)

print("S8 WORD VERSION CREATED CORRECTLY")
print("---------------------------------")
print(f"Output file: {OUTPUT_FILE}")
print(f"S8A rows: {len(a_pub)}")
print(f"S8B rows: {len(b_pub)}")
print(f"S8C rows: {len(c_pub)}")
print()
print("The audited S8A, S8B and S8C CSV files were NOT modified.")