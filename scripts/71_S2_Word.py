import pandas as pd
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


INPUT_A = Path("Supplementary_Table_S2A.csv")
INPUT_B = Path("Supplementary_Table_S2B.csv")
OUTPUT_FILE = Path("Supplementary_Table_S2_WORD.docx")


# ------------------------------------------------------------
# READ AUDITED SOURCE FILES
# ------------------------------------------------------------

if not INPUT_A.exists():
    raise FileNotFoundError(f"No se encontró: {INPUT_A.resolve()}")

if not INPUT_B.exists():
    raise FileNotFoundError(f"No se encontró: {INPUT_B.resolve()}")

df_a = pd.read_csv(INPUT_A)
df_b = pd.read_csv(INPUT_B)

if len(df_a) != 18:
    raise ValueError(
        f"S2A debería contener 18 filas, pero contiene {len(df_a)}."
    )

if len(df_b) != 12:
    raise ValueError(
        f"S2B debería contener 12 filas, pero contiene {len(df_b)}."
    )


# ------------------------------------------------------------
# WORD FORMATTING FUNCTIONS
# ------------------------------------------------------------

def shade_cell(cell, fill="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, value=70):
    tc_pr = cell._tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin in ["top", "start", "bottom", "end"]:
        element = tc_mar.find(qn(f"w:{margin}"))

        if element is None:
            element = OxmlElement(f"w:{margin}")
            tc_mar.append(element)

        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")

    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)

    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def format_table(table, widths, header_size=8, body_size=7.5):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for row in table.rows:
        prevent_row_split(row)

        for j, cell in enumerate(row.cells):
            set_cell_width(cell, widths[j])
            set_cell_margins(cell)

            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1

                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(body_size)

    for cell in table.rows[0].cells:
        shade_cell(cell)

        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(header_size)

    repeat_header(table.rows[0])


def add_panel_heading(document, text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)

    r = p.add_run(text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9)


def display_value(value):
    if pd.isna(value):
        return ""

    if isinstance(value, float):
        return f"{value:.4f}"

    return str(value)


# ------------------------------------------------------------
# CREATE DOCUMENT
# ------------------------------------------------------------

document = Document()

section = document.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = (
    section.page_height,
    section.page_width,
)

section.top_margin = Cm(1.4)
section.bottom_margin = Cm(1.4)
section.left_margin = Cm(1.3)
section.right_margin = Cm(1.3)

document.styles["Normal"].font.name = "Arial"
document.styles["Normal"].font.size = Pt(8)


# ------------------------------------------------------------
# TITLE
# ------------------------------------------------------------

p = document.add_paragraph()
p.paragraph_format.space_after = Pt(5)

r = p.add_run(
    "Supplementary Table S2. "
    "Pairwise transcriptional similarity among estradiol, "
    "calcitriol and calcipotriol across matched experimental blocks."
)

r.bold = True
r.font.name = "Arial"
r.font.size = Pt(10)


# ------------------------------------------------------------
# PANEL A
# ------------------------------------------------------------

add_panel_heading(
    document,
    "A. Descriptive pairwise transcriptional similarity"
)

# Use source columns but publication-friendly headings
panel_a = df_a.copy()

rename_a = {
    "Gene_space": "Gene space",
    "Treatment_pair": "Treatment pair",
    "Metric": "Similarity metric",
    "N_blocks": "N blocks",
    "Mean_similarity": "Mean similarity",
    "Median_similarity": "Median similarity",
    "Minimum_similarity": "Minimum similarity",
    "Maximum_similarity": "Maximum similarity",
}
panel_a = panel_a.rename(
    columns={k: v for k, v in rename_a.items() if k in panel_a.columns}
)

table_a = document.add_table(
    rows=1,
    cols=len(panel_a.columns)
)

for j, column in enumerate(panel_a.columns):
    table_a.rows[0].cells[j].text = str(column)

for _, row in panel_a.iterrows():
    cells = table_a.add_row().cells

    for j, value in enumerate(row):
        cells[j].text = display_value(value)

# Equal-ish widths determined dynamically from number of columns
usable_width_a = 26.5
widths_a = [usable_width_a / len(panel_a.columns)] * len(panel_a.columns)

format_table(
    table_a,
    widths_a,
    header_size=7.5,
    body_size=7.3
)


# ------------------------------------------------------------
# PANEL B
# ------------------------------------------------------------

add_panel_heading(
    document,
    "B. Exact paired comparisons of transcriptional similarity"
)

panel_b = df_b.copy()

rename_b = {
    "Gene_space": "Gene space",
    "Metric": "Similarity metric",
    "Contrast": "Contrast",
    "N_blocks": "N blocks",
    "Mean_similarity_VitD_pair": "Mean similarity, vitamin D pair",
    "Mean_similarity_comparator": "Mean similarity, comparator",
    "Mean_paired_difference": "Mean paired difference",
    "Median_paired_difference": "Median paired difference",
    "Blocks_VitD_pair_higher": "Blocks: vitamin D pair higher",
    "Blocks_VitD_pair_lower": "Blocks: vitamin D pair lower",
    "N_exact_sign_permutations": "Exact sign permutations",
    "Exact_two_sided_P": "Exact two-sided P",
}
panel_b = panel_b.rename(
    columns={k: v for k, v in rename_b.items() if k in panel_b.columns}
)

table_b = document.add_table(
    rows=1,
    cols=len(panel_b.columns)
)

for j, column in enumerate(panel_b.columns):
    table_b.rows[0].cells[j].text = str(column)

for _, row in panel_b.iterrows():
    cells = table_b.add_row().cells

    for j, value in enumerate(row):
        cells[j].text = display_value(value)

usable_width_b = 26.5
widths_b = [usable_width_b / len(panel_b.columns)] * len(panel_b.columns)

format_table(
    table_b,
    widths_b,
    header_size=7.5,
    body_size=7.3
)


# ------------------------------------------------------------
# NOTE
# ------------------------------------------------------------

p = document.add_paragraph()
p.paragraph_format.space_before = Pt(7)
p.paragraph_format.space_after = Pt(0)

r = p.add_run("Note: ")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(8)

r = p.add_run(
    "Pairwise transcriptional similarity was evaluated within the "
    "four matched experimental blocks using Pearson correlation, "
    "Spearman rank correlation and cosine similarity in both the "
    "Full and Landmark gene spaces (Table S2A). Table S2B reports "
    "two-sided exact paired sign-flip comparisons between "
    "calcitriol–calcipotriol similarity and each estradiol-containing "
    "pair. Exact P values were obtained from all 2⁴ = 16 possible "
    "sign configurations across the four matched blocks. Full and "
    "Landmark analyses were performed on the same underlying "
    "perturbational signatures and therefore should not be considered "
    "independent analyses."
)

r.font.name = "Arial"
r.font.size = Pt(8)


# ------------------------------------------------------------
# SAVE
# ------------------------------------------------------------

document.save(OUTPUT_FILE)

print()
print("S2 WORD VERSION CREATED CORRECTLY")
print("---------------------------------")
print(f"Output file: {OUTPUT_FILE}")
print(f"Panel A rows: {len(panel_a)}")
print(f"Panel B rows: {len(panel_b)}")
print()
print("The audited S2A and S2B CSV files were NOT modified.")