from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION


# ============================================================
# PATHS
# ============================================================

BASE = Path(".")

OUTPUT = BASE / "Supplementary_Material_COMPILED.docx"

SOURCE_FILES = [
    BASE / "Supplementary_Table_S1_WORD.docx",
    BASE / "Supplementary_Table_S2_WORD.docx",
    BASE / "Supplementary_Table_S4_WORD.docx",
    BASE / "Supplementary_Table_S5_WORD.docx",
    BASE / "Supplementary_Table_S6_WORD.docx",
    BASE / "Supplementary_Table_S7_WORD.docx",
    BASE / "Supplementary_Table_S8_WORD.docx",
    BASE / "Supplementary_Table_S9_WORD.docx",
]


# ============================================================
# CHECK SOURCE FILES
# ============================================================

missing = [str(f) for f in SOURCE_FILES if not f.exists()]

if missing:
    print("ERROR: The following source files were not found:")
    for f in missing:
        print(" -", f)
    raise SystemExit(1)


# ============================================================
# HELPERS
# ============================================================

def copy_section_format(source_section, target_section):
    """
    Copy page size, orientation and margins from a source Word section
    to a target Word section.
    """
    target_section.orientation = source_section.orientation
    target_section.page_width = source_section.page_width
    target_section.page_height = source_section.page_height

    target_section.top_margin = source_section.top_margin
    target_section.bottom_margin = source_section.bottom_margin
    target_section.left_margin = source_section.left_margin
    target_section.right_margin = source_section.right_margin

    target_section.header_distance = source_section.header_distance
    target_section.footer_distance = source_section.footer_distance


def append_document(master, source_path):
    """
    Append the body of a Word document to the master document.
    A new section is created first and its page setup is copied
    from the source document.
    """
    source = Document(source_path)

    new_section = master.add_section(WD_SECTION.NEW_PAGE)
    copy_section_format(source.sections[0], new_section)

    source_body = source.element.body

    for element in source_body.iterchildren():
        # Do not copy the final body-level section properties.
        # Page setup is handled explicitly above.
        if element.tag.endswith("sectPr"):
            continue

        master.element.body.insert(
            len(master.element.body) - 1,
            deepcopy(element)
        )


def add_heading(document, text, size=11, centred=False):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)

    if centred:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    r = p.add_run(text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(size)

    return p


def add_body_paragraph(document, text):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)

    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9)

    return p


# ============================================================
# CREATE MASTER DOCUMENT
# ============================================================

doc = Document()

# Initial section: portrait for Supplementary Methods
section = doc.sections[0]

section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(9)


# ============================================================
# TITLE
# ============================================================

add_heading(
    doc,
    "Supplementary Material",
    size=14,
    centred=True
)

add_heading(
    doc,
    "Supplementary Methods",
    size=11,
    centred=False
)


# ============================================================
# SUPPLEMENTARY METHODS S1
# ============================================================

add_heading(
    doc,
    "S1. Leading-edge analysis of the robust convergent Hallmark programmes",
    size=10
)

add_body_paragraph(
    doc,
    "A descriptive leading-edge analysis was performed for the three Hallmark "
    "programmes that showed significant concordant negative enrichment across "
    "estradiol, calcitriol and calcipotriol in both the Full and Landmark "
    "analyses: HALLMARK_E2F_TARGETS, HALLMARK_G2M_CHECKPOINT and "
    "HALLMARK_MYC_TARGETS_V1. Treatment-specific gene rankings and Hallmark "
    "gene-set definitions were the same as those used for the preranked "
    "enrichment analysis described in Section 2.7 of the main manuscript."
)

add_body_paragraph(
    doc,
    "For each treatment, Hallmark programme and gene space, genes were ranked "
    "in descending order of treatment-level mean MODZ. The weighted running-sum "
    "statistic was recalculated using a weighting exponent of 1, with hit "
    "weights proportional to the absolute mean MODZ value and a constant "
    "penalty for genes outside the tested Hallmark set. The enrichment score "
    "(ES) was defined as the maximum deviation from zero with the greatest "
    "absolute magnitude. For negatively enriched programmes, the leading edge "
    "was defined as the Hallmark genes occurring from the rank position of the "
    "minimum running-sum value to the end of the ranked list; for positively "
    "enriched programmes, it was defined as the Hallmark genes occurring from "
    "the beginning of the ranked list through the position of the maximum "
    "running-sum value."
)

add_body_paragraph(
    doc,
    "The procedure was performed separately in the Full and Landmark gene "
    "spaces. Within each Hallmark programme and gene space, leading-edge genes "
    "shared by estradiol, calcitriol and calcipotriol were identified by "
    "intersection of the three treatment-specific leading-edge gene sets. "
    "Genes belonging to the three-treatment shared leading edge of at least "
    "two of the three robust Hallmark programmes were additionally classified "
    "as recurrent genes of the convergent transcriptional core. These analyses "
    "were descriptive and no additional statistical significance testing was "
    "performed."
)


# ============================================================
# SUPPLEMENTARY METHODS S2
# ============================================================

add_heading(
    doc,
    "S2. Receptor-axis transcriptional analysis",
    size=10
)

add_body_paragraph(
    doc,
    "A descriptive receptor-axis analysis was performed to examine "
    "perturbation-associated transcriptional responses of genes related to "
    "vitamin D and estrogen signalling within the primary matched dataset. "
    "Nine genes were predefined for analysis: VDR, RXRA, RXRB, RXRG, CYP24A1 "
    "and CYP27B1 for the vitamin D-related axis, and ESR1, ESR2 and GPER1 for "
    "the estrogen-related axis. Gene symbols were matched directly to the "
    "L1000 gene annotation, and Landmark status was determined from the "
    "pr_is_lm field of the GSE92742 gene metadata."
)

add_body_paragraph(
    doc,
    "For each receptor-axis gene, Level 5 MODZ values were extracted from each "
    "of the 12 primary signatures and examined across the four matched "
    "experimental contexts and three treatments. For each treatment–gene "
    "combination, the mean, median, minimum and maximum MODZ values were "
    "calculated, together with the numbers of contexts showing positive and "
    "negative MODZ values. Directional consistency was classified as uniformly "
    "positive or negative when all four contexts showed the same sign, as "
    "predominantly positive or negative when three of four contexts showed the "
    "same sign, and as heterogeneous otherwise."
)

add_body_paragraph(
    doc,
    "Context-specific comparisons were additionally generated for estradiol, "
    "calcitriol and calcipotriol. For descriptive comparison with estradiol, "
    "the mean MODZ response to the two vitamin D perturbations was calculated "
    "within each context, and the difference between the estradiol MODZ value "
    "and this vitamin D mean was determined for each receptor-axis gene. No "
    "inferential statistical testing was performed for these receptor-axis "
    "analyses."
)

add_body_paragraph(
    doc,
    "Because L1000 Level 5 MODZ values represent differential transcriptional "
    "responses to perturbation, receptor-axis MODZ values were not interpreted "
    "as measures of basal receptor expression or as direct measures of receptor "
    "functional activity. Genes not annotated as Landmark transcripts were "
    "explicitly considered computationally inferred L1000 measurements."
)


# ============================================================
# SUPPLEMENTARY TABLES
# ============================================================

doc.add_page_break()

add_heading(
    doc,
    "Supplementary Tables",
    size=12,
    centred=False
)


# S1
append_document(
    doc,
    BASE / "Supplementary_Table_S1_WORD.docx"
)

# S2
append_document(
    doc,
    BASE / "Supplementary_Table_S2_WORD.docx"
)


# ============================================================
# S3 — ELECTRONIC DATA TABLE
# ============================================================

s3_section = doc.add_section(WD_SECTION.NEW_PAGE)

s3_section.page_width = Cm(21.0)
s3_section.page_height = Cm(29.7)
s3_section.top_margin = Cm(2.0)
s3_section.bottom_margin = Cm(2.0)
s3_section.left_margin = Cm(2.0)
s3_section.right_margin = Cm(2.0)

add_heading(
    doc,
    "Supplementary Data Table S3. Genes showing directionally consistent "
    "transcriptional responses across all four matched experimental contexts.",
    size=10
)

add_body_paragraph(
    doc,
    "Note: Genes were classified as directionally consistent when their Level "
    "5 MODZ values showed the same sign across all four matched experimental "
    "contexts for a given treatment. Positive and negative consistency were "
    "evaluated separately for estradiol, calcitriol and calcipotriol. Mean, "
    "median and mean absolute MODZ values summarise the transcriptional "
    "response across the four contexts. Landmark status indicates whether a "
    "gene was directly measured by the L1000 assay; non-Landmark genes "
    "correspond to computationally inferred measurements. The 4/4 "
    "directional-consistency criterion is descriptive and should not be "
    "interpreted as a statistical significance test."
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)

r = p.add_run(
    "Supplementary Data Table S3 is provided as a separate electronic CSV file "
    "(Supplementary_Data_Table_S3_PUBLICATION.csv)."
)
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(9)


# ============================================================
# S4–S9
# ============================================================

for source_file in [
    BASE / "Supplementary_Table_S4_WORD.docx",
    BASE / "Supplementary_Table_S5_WORD.docx",
    BASE / "Supplementary_Table_S6_WORD.docx",
    BASE / "Supplementary_Table_S7_WORD.docx",
    BASE / "Supplementary_Table_S8_WORD.docx",
    BASE / "Supplementary_Table_S9_WORD.docx",
]:
    append_document(doc, source_file)


# ============================================================
# SAVE
# ============================================================

doc.save(OUTPUT)

print("SUPPLEMENTARY MATERIAL COMPILED CORRECTLY")
print("-----------------------------------------")
print(f"Output file: {OUTPUT.name}")
print()
print("Included:")
print(" - Supplementary Methods S1")
print(" - Supplementary Methods S2")
print(" - Supplementary Table S1")
print(" - Supplementary Table S2")
print(" - Supplementary Data Table S3 reference")
print(" - Supplementary Tables S4-S9")
print()
print("All source Word and CSV files were left unchanged.")