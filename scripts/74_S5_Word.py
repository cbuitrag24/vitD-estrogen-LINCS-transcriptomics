import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

INPUT_FILE = "Supplementary_Table_S5.csv"
OUTPUT_FILE = "Supplementary_Table_S5_WORD.docx"

TITLE = (
    "Supplementary Table S5. Hallmark gene-set enrichment results for "
    "estradiol, calcitriol and calcipotriol in the Landmark gene space."
)

NOTE = (
    "Note: Preranked Hallmark gene-set enrichment analysis was repeated in the "
    "Landmark gene space as a robustness analysis restricted to directly measured "
    "L1000 transcripts. Hallmark gene sets containing at least 15 Landmark genes "
    "were retained, resulting in 35 evaluable programmes. Enrichment scores (ES), "
    "normalised enrichment scores (NES), empirical P values based on 5,000 randomly "
    "sampled gene sets of identical size, and Benjamini–Hochberg-adjusted empirical "
    "P values were calculated using the same procedure as in the Full analysis. "
    "Multiple-testing correction was applied separately within each treatment. "
    "Hallmark programmes with BH-adjusted empirical P < 0.05 were considered "
    "significantly enriched. Because the Full and Landmark analyses derive from "
    "the same underlying perturbational signatures, the Landmark analysis was "
    "considered a robustness analysis rather than an independent validation."
)

# Read the audited source table
df = pd.read_csv(INPUT_FILE)

# Safety check
if len(df) != 105:
    raise ValueError(
        f"Unexpected number of rows in S5: {len(df)}. Expected 105."
    )

print("Source columns:")
for col in df.columns:
    print(f"  - {col}")

# Publication-friendly headings.
# Only columns actually present in the audited CSV will be renamed.
rename_columns = {
    "Treatment": "Treatment",
    "Hallmark_programme": "Hallmark programme",
    "N_genes_in_ranking": "Genes in ranking",
    "N_Landmark_genes": "Landmark genes",
    "N_genes": "Gene-set size",
    "ES": "ES",
    "NES": "NES",
    "Empirical_P": "Empirical P",
    "BH_adjusted_empirical_P": "BH-adjusted empirical P",
    "BH_significant_0.05": "BH significant (P < 0.05)",
}

publication = df.rename(
    columns={
        key: value
        for key, value in rename_columns.items()
        if key in df.columns
    }
).copy()


def display_value(value):
    if pd.isna(value):
        return ""

    if isinstance(value, float):
        if abs(value) < 0.0001 and value != 0:
            return f"{value:.3e}"
        return f"{value:.4f}"

    return str(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


document = Document()

# Landscape page
section = document.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(1.3)
section.bottom_margin = Cm(1.3)
section.left_margin = Cm(1.3)
section.right_margin = Cm(1.3)

# Default font
styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(8)

# Title
title_p = document.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

title_run = title_p.add_run(TITLE)
title_run.bold = True
title_run.font.name = "Arial"
title_run.font.size = Pt(10)

# Table
table = document.add_table(
    rows=1,
    cols=len(publication.columns)
)

table.style = "Table Grid"
table.autofit = True

# Header
header = table.rows[0]
set_repeat_table_header(header)

for j, col in enumerate(publication.columns):
    cell = header.cells[j]
    cell.text = str(col)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "D9E2F3")

    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for run in paragraph.runs:
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(7)

# Data
for _, row in publication.iterrows():
    cells = table.add_row().cells

    for j, value in enumerate(row):
        cells[j].text = display_value(value)
        cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for paragraph in cells[j].paragraphs:
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(7)

# Note
note_p = document.add_paragraph()
note_run = note_p.add_run(NOTE)
note_run.font.name = "Arial"
note_run.font.size = Pt(8)

# Save
document.save(OUTPUT_FILE)

print()
print("S5 WORD VERSION CREATED CORRECTLY")
print("---------------------------------")
print(f"Output file: {OUTPUT_FILE}")
print(f"Rows: {len(publication)}")
print(f"Columns: {len(publication.columns)}")
print()
print("Publication headings:")
for col in publication.columns:
    print(f"  - {col}")
print()
print("The audited Supplementary_Table_S5.csv file was NOT modified.")