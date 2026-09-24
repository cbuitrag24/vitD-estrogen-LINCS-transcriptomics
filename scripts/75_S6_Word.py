import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FILE_A = "Supplementary_Table_S6A.csv"
FILE_B = "Supplementary_Table_S6B.csv"
FILE_C = "Supplementary_Table_S6C.csv"

OUTPUT_FILE = "Supplementary_Table_S6_WORD.docx"

TITLE = (
    "Supplementary Table S6. Leading-edge analysis of the robust convergent "
    "Hallmark programmes across estradiol, calcitriol and calcipotriol."
)

NOTE = (
    "Note: Leading-edge analysis was performed for the three Hallmark programmes "
    "showing significant concordant negative enrichment across estradiol, "
    "calcitriol and calcipotriol in both the Full and Landmark analyses: "
    "HALLMARK_E2F_TARGETS, HALLMARK_G2M_CHECKPOINT and "
    "HALLMARK_MYC_TARGETS_V1. Table S6A summarises treatment-specific leading "
    "edges in each gene space. Table S6B reports leading-edge genes shared by "
    "all three treatments within each Hallmark programme, and Table S6C "
    "identifies genes occurring in the three-treatment shared leading edge of "
    "at least two of the three robust Hallmark programmes. The analysis was "
    "performed separately in the Full and Landmark gene spaces and was "
    "descriptive; no additional statistical significance testing was applied."
)

# ------------------------------------------------------------
# Read audited source files
# ------------------------------------------------------------

a = pd.read_csv(FILE_A)
b = pd.read_csv(FILE_B)
c = pd.read_csv(FILE_C)

# Safety checks
if len(a) != 18:
    raise ValueError(f"Unexpected S6A row count: {len(a)}. Expected 18.")

if len(b) != 168:
    raise ValueError(f"Unexpected S6B row count: {len(b)}. Expected 168.")

if len(c) != 46:
    raise ValueError(f"Unexpected S6C row count: {len(c)}. Expected 46.")

expected_a = [
    "Treatment",
    "Gene_space",
    "Hallmark_programme",
    "ES",
    "ES_direction",
    "Peak_rank",
    "N_Hallmark_genes_in_ranking",
    "N_leading_edge_genes",
]

expected_b = [
    "Gene_space",
    "Hallmark_programme",
    "Gene",
]

expected_c = [
    "Gene_space",
    "Gene",
    "N_robust_Hallmark_programmes",
]

if list(a.columns) != expected_a:
    raise ValueError("Unexpected columns in S6A.")

if list(b.columns) != expected_b:
    raise ValueError("Unexpected columns in S6B.")

if list(c.columns) != expected_c:
    raise ValueError("Unexpected columns in S6C.")

# ------------------------------------------------------------
# Publication copies
# ------------------------------------------------------------

a_pub = a.rename(
    columns={
        "Treatment": "Treatment",
        "Gene_space": "Gene space",
        "Hallmark_programme": "Hallmark programme",
        "ES": "ES",
        "ES_direction": "ES direction",
        "Peak_rank": "Peak rank",
        "N_Hallmark_genes_in_ranking": "Hallmark genes in ranking",
        "N_leading_edge_genes": "Leading-edge genes",
    }
).copy()

# Editorial translation only in the Word copy
a_pub["ES direction"] = (
    a_pub["ES direction"]
    .replace(
        {
            "negativo": "Negative",
            "positivo": "Positive",
            "negative": "Negative",
            "positive": "Positive",
        }
    )
)

b_pub = b.rename(
    columns={
        "Gene_space": "Gene space",
        "Hallmark_programme": "Hallmark programme",
        "Gene": "Gene",
    }
).copy()

c_pub = c.rename(
    columns={
        "Gene_space": "Gene space",
        "Gene": "Gene",
        "N_robust_Hallmark_programmes": "Robust Hallmark programmes",
    }
).copy()

# ------------------------------------------------------------
# Formatting functions
# ------------------------------------------------------------

def display_value(value):
    if pd.isna(value):
        return ""

    if isinstance(value, float):
        if abs(value) < 0.0001 and value != 0:
            return f"{value:.3e}"
        return f"{value:.4f}"

    return str(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_panel_heading(document, text):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)

    r = p.add_run(text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9)


def add_table(document, dataframe):
    table = document.add_table(
        rows=1,
        cols=len(dataframe.columns)
    )

    table.style = "Table Grid"
    table.autofit = True

    header = table.rows[0]
    set_repeat_table_header(header)

    for j, col in enumerate(dataframe.columns):
        cell = header.cells[j]
        cell.text = str(col)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9E2F3")

        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(7)

    for _, row in dataframe.iterrows():
        cells = table.add_row().cells

        for j, value in enumerate(row):
            cells[j].text = display_value(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for paragraph in cells[j].paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7)

    return table


# ------------------------------------------------------------
# Create Word document
# ------------------------------------------------------------

document = Document()

section = document.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(1.3)
section.bottom_margin = Cm(1.3)
section.left_margin = Cm(1.3)
section.right_margin = Cm(1.3)

document.styles["Normal"].font.name = "Arial"
document.styles["Normal"].font.size = Pt(8)

# Main title
title_p = document.add_paragraph()
title_run = title_p.add_run(TITLE)
title_run.bold = True
title_run.font.name = "Arial"
title_run.font.size = Pt(10)

# S6A
add_panel_heading(
    document,
    "S6A. Treatment-specific leading-edge characteristics"
)
add_table(document, a_pub)

# Page break before S6B
document.add_page_break()

# S6B
add_panel_heading(
    document,
    "S6B. Leading-edge genes shared by all three treatments"
)
add_table(document, b_pub)

# Page break before S6C
document.add_page_break()

# S6C
add_panel_heading(
    document,
    "S6C. Recurrent genes across robust Hallmark programmes"
)
add_table(document, c_pub)

# Note
note_p = document.add_paragraph()
note_p.paragraph_format.space_before = Pt(6)

note_run = note_p.add_run(NOTE)
note_run.font.name = "Arial"
note_run.font.size = Pt(8)

# Save
document.save(OUTPUT_FILE)

# ------------------------------------------------------------
# Console confirmation
# ------------------------------------------------------------

print("S6 WORD VERSION CREATED CORRECTLY")
print("---------------------------------")
print(f"Output file: {OUTPUT_FILE}")
print(f"S6A rows: {len(a_pub)}")
print(f"S6B rows: {len(b_pub)}")
print(f"S6C rows: {len(c_pub)}")
print()
print("S6A ES directions:")
print(a_pub["ES direction"].value_counts(dropna=False).to_dict())
print()
print(
    "The audited S6A, S6B and S6C CSV files were NOT modified."
)